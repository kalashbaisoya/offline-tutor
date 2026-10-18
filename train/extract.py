"""S0 + S1: inventory and text extraction for books / slides / notes.

Walks a folder of PDF / PPTX / TXT / MD files and writes:
  <out>/corpus.jsonl    one record per page or slide: {doc, kind, page, chars, text}
  <out>/inventory.json  per document: pages, total chars, pages that look scanned (little or no text layer),
                        and whether Devanagari (Hindi) text is present.

Scanned pages are only flagged here; OCR is a separate, later step so this one stays fast and dependency-light.

Usage: python train/extract.py RAW_DIR OUT_DIR
"""

import json
import re
import sys
from pathlib import Path

DEVANAGARI = re.compile(r"[ऀ-ॿ]")
SCANNED_CHARS = 200  # a page with fewer extracted characters than this is probably an image scan


def clean(text):
    text = text.replace("\r", "\n").replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def pdf_pages(path):
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            yield i + 1, page.get_textpage().get_text_range()
    finally:
        pdf.close()


def pptx_slides(path):
    from pptx import Presentation

    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        for shape in slide.shapes:
            if shape.has_table:
                parts += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide:
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        yield i, "\n".join(parts)


def text_file(path):
    yield 1, path.read_text(encoding="utf-8", errors="replace")


READERS = {".pdf": ("pdf", pdf_pages), ".pptx": ("pptx", pptx_slides), ".txt": ("text", text_file),
           ".md": ("text", text_file)}


def main(raw_dir, out_dir):
    raw, out = Path(raw_dir), Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    inventory, skipped = [], []
    with open(out / "corpus.jsonl", "w", encoding="utf-8") as corpus:
        for path in sorted(p for p in raw.rglob("*") if p.is_file()):
            reader = READERS.get(path.suffix.lower())
            if not reader:
                skipped.append(str(path.relative_to(raw)))
                continue
            kind, read = reader
            doc = str(path.relative_to(raw))
            pages = chars = 0
            scanned, hindi = [], False
            try:
                for page, text in read(path):
                    text = clean(text)
                    pages += 1
                    chars += len(text)
                    if kind == "pdf" and len(text) < SCANNED_CHARS:
                        scanned.append(page)
                    hindi = hindi or bool(DEVANAGARI.search(text))
                    corpus.write(json.dumps(dict(doc=doc, kind=kind, page=page, chars=len(text), text=text),
                                            ensure_ascii=False) + "\n")
            except Exception as e:  # a corrupt file should not stop the whole run
                inventory.append(dict(doc=doc, kind=kind, error=f"{type(e).__name__}: {e}"))
                continue
            inventory.append(dict(doc=doc, kind=kind, pages=pages, chars=chars, scanned_pages=len(scanned),
                                  scanned_first=scanned[:20], has_devanagari=hindi))
    summary = dict(documents=len(inventory), pages=sum(d.get("pages", 0) for d in inventory),
                   chars=sum(d.get("chars", 0) for d in inventory),
                   scanned_pages=sum(d.get("scanned_pages", 0) for d in inventory),
                   errors=[d for d in inventory if "error" in d], skipped_files=skipped)
    (out / "inventory.json").write_text(json.dumps(dict(summary=summary, documents=inventory), indent=2,
                                                   ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main(*sys.argv[1:3])
